"""Воспроизводимая генерация синтетических демо-данных DataStory AI.

Все данные полностью вымышлены: платёжная система «DemoPay KZ» не существует.
Генерация детерминирована (фиксированный seed), повторный запуск даёт те же значения.

Создаёт в каталоге data/demo:
  transactions_2026.xlsx        статистика январь–июнь 2026 (месяц x канал)
  business_metrics.pdf          описание показателей (кириллица, встроенный шрифт)
  transactions_screenshot.png   изображение части таблицы для Vision-моделей
  expected_metrics.json         контрольные показатели для автотестов

Запуск:  python scripts/generate_demo_data.py [--out data/demo]
"""
from __future__ import annotations

import argparse
import io
import json
import random
import re
import sys
import zipfile
from datetime import datetime
from pathlib import Path

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw, ImageFont

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_OUT = PROJECT_ROOT / "data" / "demo"

SEED = 2026
MONTHS = ["2026-01", "2026-02", "2026-03", "2026-04", "2026-05", "2026-06"]
COLUMNS = ["Month", "Transactions", "Successful", "Failed", "Amount_KZT", "Channel"]
DROP_MONTH = "2026-03"
SCREENSHOT_MONTHS = ["2026-02", "2026-03"]
FIXED_DATE = datetime(2026, 7, 1, 0, 0, 0)

# Базовые параметры каналов: число транзакций в месяц, успешность, средняя сумма (KZT).
CHANNELS = {
    "Mobile": {"tx": 182_000, "success": 0.978, "avg": 8_400, "march_drop": 0.055},
    "Web": {"tx": 96_000, "success": 0.972, "avg": 14_200, "march_drop": 0.065},
    "API": {"tx": 41_000, "success": 0.986, "avg": 31_500, "march_drop": 0.035},
}
MONTHLY_GROWTH = 0.02


# --------------------------------------------------------------------------- данные
def build_rows() -> list[dict]:
    """Строки датасета: 6 месяцев x 3 канала. Порядок вызовов rng фиксирован."""
    rng = random.Random(SEED)
    rows: list[dict] = []
    for index, month in enumerate(MONTHS):
        trend = (1 + MONTHLY_GROWTH) ** index
        for channel, cfg in CHANNELS.items():
            transactions = round(cfg["tx"] * trend * (1 + rng.uniform(-0.03, 0.03)))
            success_rate = cfg["success"] + rng.uniform(-0.0025, 0.0025)
            if month == DROP_MONTH:
                success_rate -= cfg["march_drop"] * (1 + rng.uniform(-0.1, 0.1))
            successful = round(transactions * success_rate)
            avg_amount = cfg["avg"] * (1 + rng.uniform(-0.025, 0.025))
            rows.append(
                {
                    "Month": month,
                    "Transactions": transactions,
                    "Successful": successful,
                    "Failed": transactions - successful,
                    "Amount_KZT": round(transactions * avg_amount),
                    "Channel": channel,
                }
            )
    return rows


def validate_rows(rows: list[dict]) -> None:
    assert len(rows) == len(MONTHS) * len(CHANNELS)
    for row in rows:
        assert set(row) == set(COLUMNS), row
        assert all(v not in (None, "") for v in row.values()), row
        assert row["Transactions"] == row["Successful"] + row["Failed"], row
        assert row["Failed"] > 0 and row["Successful"] > 0 and row["Amount_KZT"] > 0, row
    pairs = {(r["Month"], r["Channel"]) for r in rows}
    assert len(pairs) == len(rows), "дубликаты пары месяц/канал"


# --------------------------------------------------------------------------- контрольные показатели
def _agg(rows: list[dict]) -> dict:
    tx = sum(r["Transactions"] for r in rows)
    ok = sum(r["Successful"] for r in rows)
    failed = sum(r["Failed"] for r in rows)
    amount = sum(r["Amount_KZT"] for r in rows)
    return {
        "transactions": tx,
        "successful": ok,
        "failed": failed,
        "amount_kzt": amount,
        "success_rate_pct": round(ok / tx * 100, 4),
        "avg_transaction_amount_kzt": round(amount / tx, 4),
    }


def compute_expected(rows: list[dict]) -> dict:
    """Ожидаемые показатели, посчитанные обычным Python (без pandas)."""
    by_month = {m: _agg([r for r in rows if r["Month"] == m]) for m in MONTHS}
    by_channel = {c: _agg([r for r in rows if r["Channel"] == c]) for c in CHANNELS}
    by_month_channel = [
        {"month": r["Month"], "channel": r["Channel"], **_agg([r])} for r in rows
    ]

    lowest = min(by_month.items(), key=lambda kv: kv[1]["success_rate_pct"])
    other_months = [r for r in rows if r["Month"] != DROP_MONTH]
    baseline = _agg(other_months)["success_rate_pct"]
    february = by_month["2026-02"]["success_rate_pct"]
    drop = by_month[DROP_MONTH]["success_rate_pct"]

    shot_rows = [r for r in rows if r["Month"] in SCREENSHOT_MONTHS]
    return {
        "meta": {
            "description": "Ожидаемые контрольные показатели для синтетического датасета DemoPay KZ.",
            "synthetic": True,
            "seed": SEED,
            "generator": "scripts/generate_demo_data.py",
            "formulas": {
                "success_rate_pct": "Successful / Transactions * 100 (при агрегации: суммы, затем деление)",
                "avg_transaction_amount_kzt": "Amount_KZT / Transactions (при агрегации: суммы, затем деление)",
            },
        },
        "row_count": len(rows),
        "columns": COLUMNS,
        "months": MONTHS,
        "channels": list(CHANNELS),
        "totals": _agg(rows),
        "by_month": by_month,
        "by_channel": by_channel,
        "by_month_channel": by_month_channel,
        "anomaly": {
            "lowest_success_rate_month": lowest[0],
            "lowest_success_rate_pct": lowest[1]["success_rate_pct"],
            "march_success_rate_pct": drop,
            "february_success_rate_pct": february,
            "baseline_success_rate_excl_march_pct": baseline,
            "march_vs_february_pp": round(drop - february, 4),
            "march_vs_baseline_pp": round(drop - baseline, 4),
            "march_vs_february_pp_by_channel": {
                c: round(
                    _agg([r for r in rows if r["Month"] == DROP_MONTH and r["Channel"] == c])[
                        "success_rate_pct"
                    ]
                    - _agg([r for r in rows if r["Month"] == "2026-02" and r["Channel"] == c])[
                        "success_rate_pct"
                    ],
                    4,
                )
                for c in CHANNELS
            },
        },
        "leaders": {
            "highest_transactions_channel": max(by_channel, key=lambda c: by_channel[c]["transactions"]),
            "highest_amount_channel": max(by_channel, key=lambda c: by_channel[c]["amount_kzt"]),
            "highest_avg_amount_channel": max(
                by_channel, key=lambda c: by_channel[c]["avg_transaction_amount_kzt"]
            ),
            "highest_success_rate_channel": max(by_channel, key=lambda c: by_channel[c]["success_rate_pct"]),
        },
        "screenshot": {
            "file": "transactions_screenshot.png",
            "months": SCREENSHOT_MONTHS,
            "row_count": len(shot_rows),
            "rows": shot_rows,
            "totals": _agg(shot_rows),
        },
    }


# --------------------------------------------------------------------------- Excel
def write_xlsx(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "transactions"
    ws.append(COLUMNS)
    for row in rows:
        ws.append([row[c] for c in COLUMNS])

    header_fill = PatternFill("solid", fgColor="4F46E5")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for col_cells in ws.iter_cols(min_row=2, min_col=2, max_col=5):
        for cell in col_cells:
            cell.number_format = "0"
    for letter, width in zip("ABCDEF", (10, 14, 14, 12, 16, 10)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"

    wb.properties.creator = "DataStory AI demo generator"
    wb.properties.created = wb.properties.modified = FIXED_DATE
    wb.save(path)
    _normalize_zip(path)


def _normalize_zip(path: Path) -> None:
    """xlsx — это ZIP: фиксируем время записей и дату modified (openpyxl ставит «сейчас»),
    чтобы файл был побайтно воспроизводим."""
    with zipfile.ZipFile(path) as src:
        entries = [(info.filename, src.read(info.filename)) for info in src.infolist()]
    fixed = FIXED_DATE.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    entries = [
        (name, re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>" + fixed, data) if name == "docProps/core.xml" else data)
        for name, data in entries
    ]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for name, data in entries:
            info = zipfile.ZipInfo(name, date_time=(2026, 7, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            dst.writestr(info, data)


# --------------------------------------------------------------------------- шрифты
def _fonts() -> tuple[pymupdf.Font, pymupdf.Font, bytes, bytes]:
    """Встроенный Nimbus Sans из PyMuPDF (поддерживает кириллицу) — без системных шрифтов."""
    reg_buf, bold_buf = pymupdf.Font("helv").buffer, pymupdf.Font("hebo").buffer
    return pymupdf.Font(fontbuffer=reg_buf), pymupdf.Font(fontbuffer=bold_buf), reg_buf, bold_buf


# --------------------------------------------------------------------------- PDF
PDF_TITLE = "DemoPay KZ — описание бизнес-показателей"
PDF_INTRO = (
    "Документ создан для демонстрации DataStory AI. Платёжная система, данные и события в нём "
    "полностью вымышлены; совпадения с реальными организациями случайны."
)

# (заголовок, [абзацы]); абзац, начинающийся с «• », выводится как элемент списка.
PDF_SECTIONS: list[tuple[str, list[str]]] = [
    (
        "1. Описание платёжной системы",
        [
            "DemoPay KZ — вымышленная платёжная система, принимающая платежи в тенге (KZT) через три "
            "канала: Mobile, Web и API. Датасет transactions_2026.xlsx содержит агрегированную "
            "статистику за январь — июнь 2026 года: для каждого месяца приведено по одной строке на "
            "канал, всего 18 строк.",
        ],
    ),
    (
        "2. Описание колонок исходного датасета",
        [
            "• Month — календарный месяц в формате ГГГГ-ММ, например 2026-03.",
            "• Transactions — общее число инициированных транзакций за месяц в данном канале.",
            "• Successful — число успешно завершённых транзакций.",
            "• Failed — число неуспешных транзакций (отклонённых или завершившихся ошибкой).",
            "• Amount_KZT — общая сумма всех инициированных транзакций за месяц в данном канале, в тенге.",
            "• Channel — канал приёма платежа: Mobile, Web или API.",
            "Для каждой строки выполняется равенство: Transactions = Successful + Failed.",
        ],
    ),
    (
        "3. Определение показателя Success Rate",
        [
            "Success Rate (успешность транзакций) — доля успешно завершённых транзакций от общего "
            "числа инициированных транзакций, выраженная в процентах.",
        ],
    ),
    (
        "4. Формула Success Rate",
        [
            "Success Rate = Successful / Transactions × 100",
            "При объединении нескольких строк (месяцы, каналы) сначала суммируют Successful и "
            "Transactions, затем делят одну сумму на другую. Средний арифметический процент по строкам "
            "не используется, так как он не учитывает разный объём каналов.",
        ],
    ),
    (
        "5. Определение Transaction Volume",
        [
            "Transaction Volume (объём транзакций) — денежный объём платежей: сумма Amount_KZT за "
            "выбранный период или канал, в тенге. Число транзакций (колонка Transactions) называется "
            "«количество транзакций» и с объёмом не смешивается.",
        ],
    ),
    (
        "6. Определение Average Transaction Amount",
        [
            "Average Transaction Amount (средняя сумма транзакции) — средняя сумма одной "
            "инициированной транзакции в тенге.",
        ],
    ),
    (
        "7. Формула Average Transaction Amount",
        [
            "Average Transaction Amount = Amount_KZT / Transactions",
            "При объединении строк суммируют Amount_KZT и Transactions, затем делят.",
        ],
    ),
    (
        "8. Описание каналов",
        [
            "• Mobile — платежи из мобильного приложения. Самый массовый канал с небольшой средней суммой.",
            "• Web — платежи на сайтах и в веб-версии сервиса. Средняя сумма выше, чем в Mobile.",
            "• API — серверные интеграции партнёров. Меньше транзакций, но самые крупные суммы.",
        ],
    ),
    (
        "9. Контекстное событие",
        [
            "В марте 2026 года проводилось плановое обновление инфраструктуры обработки транзакций.",
            "Важно: наличие этого события не доказывает причинную связь со снижением успешности "
            "транзакций. Совпадение по времени — лишь гипотеза, которую необходимо проверять. Для "
            "выводов о причинах требуются дополнительные данные: журналы ошибок, коды отказов, "
            "сроки и состав работ. В аналитических выводах следует писать «совпадает по времени», "
            "а не «вызвано обновлением».",
        ],
    ),
]


class _PdfWriter:
    MARGIN = 56
    BODY, HEAD, TITLE = 10.5, 13.0, 20.0

    def __init__(self, doc: pymupdf.Document, regular: pymupdf.Font, bold: pymupdf.Font):
        self.doc, self.regular, self.bold = doc, regular, bold
        self.page: pymupdf.Page | None = None
        self.y = 0.0
        self._new_page()

    def _new_page(self) -> None:
        self.page = self.doc.new_page(width=595, height=842)  # A4
        self.y = self.MARGIN

    @property
    def width(self) -> float:
        return 595 - 2 * self.MARGIN

    def _wrap(self, text: str, font: pymupdf.Font, size: float, width: float) -> list[str]:
        lines, current = [], ""
        for word in text.split():
            candidate = f"{current} {word}".strip()
            if font.text_length(candidate, fontsize=size) <= width:
                current = candidate
            else:
                lines.append(current)
                current = word
        return lines + ([current] if current else [])

    def text(self, text: str, *, font: pymupdf.Font, size: float, indent: float = 0.0,
             gap_after: float = 6.0, color=(0.07, 0.09, 0.15), first_line_prefix: str = "") -> None:
        leading = size * 1.4
        for i, line in enumerate(self._wrap(text, font, size, self.width - indent)):
            if self.y + leading > 842 - self.MARGIN:
                self._new_page()
            writer = pymupdf.TextWriter(self.page.rect, color=color)
            x = self.MARGIN + indent
            if i == 0 and first_line_prefix:
                writer.append((self.MARGIN + indent - 12, self.y + size), first_line_prefix, font=font, fontsize=size)
            writer.append((x, self.y + size), line, font=font, fontsize=size)
            writer.write_text(self.page)
            self.y += leading
        self.y += gap_after

    def heading(self, text: str) -> None:
        if self.y + 60 > 842 - self.MARGIN:  # заголовок не остаётся один внизу страницы
            self._new_page()
        self.y += 6
        self.text(text, font=self.bold, size=self.HEAD, gap_after=4, color=(0.31, 0.27, 0.9))

    def paragraph(self, text: str) -> None:
        if text.startswith("• "):
            self.text(text[2:], font=self.regular, size=self.BODY, indent=14, gap_after=3, first_line_prefix="•")
        elif "=" in text and len(text) < 80:  # формула — отдельным акцентным блоком
            self.text(text, font=self.bold, size=11.5, indent=14, gap_after=6)
        else:
            self.text(text, font=self.regular, size=self.BODY, gap_after=6)


def write_pdf(path: Path) -> None:
    regular, bold, _, _ = _fonts()
    doc = pymupdf.open()
    pdf = _PdfWriter(doc, regular, bold)
    pdf.text(PDF_TITLE, font=bold, size=pdf.TITLE, gap_after=6)
    pdf.text(PDF_INTRO, font=regular, size=9.5, gap_after=10, color=(0.42, 0.45, 0.5))
    for title, paragraphs in PDF_SECTIONS:
        pdf.heading(title)
        for paragraph in paragraphs:
            pdf.paragraph(paragraph)

    doc.set_metadata(
        {
            "title": PDF_TITLE,
            "author": "DataStory AI demo generator",
            "subject": "Синтетические демонстрационные данные",
            "creator": "scripts/generate_demo_data.py",
            "producer": "PyMuPDF",
            "creationDate": "D:20260701000000Z",
            "modDate": "D:20260701000000Z",
        }
    )
    doc.save(path, garbage=4, deflate=True, no_new_id=True)
    doc.close()


# --------------------------------------------------------------------------- PNG
def write_png(rows: list[dict], path: Path) -> None:
    """Небольшая таблица высокого контраста, пригодная для распознавания Vision-моделью."""
    _, _, reg_buf, bold_buf = _fonts()
    font = ImageFont.truetype(io.BytesIO(reg_buf), 32)
    font_bold = ImageFont.truetype(io.BytesIO(bold_buf), 32)

    shown = [r for r in rows if r["Month"] in SCREENSHOT_MONTHS]
    table = [COLUMNS] + [[str(r[c]) for c in COLUMNS] for r in shown]
    numeric = {"Transactions", "Successful", "Failed", "Amount_KZT"}

    pad_x, row_h, margin = 24, 64, 32
    widths = [
        max(int(font_bold.getlength(line[i]) if n == 0 else font.getlength(line[i]))
            for n, line in enumerate(table)) + 2 * pad_x
        for i in range(len(COLUMNS))
    ]
    width = sum(widths) + 2 * margin
    height = row_h * len(table) + 2 * margin
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)

    grid, ink = "#374151", "#111827"
    top = margin
    for r, line in enumerate(table):
        y0, y1 = top + r * row_h, top + (r + 1) * row_h
        if r == 0:
            draw.rectangle([margin, y0, width - margin, y1], fill="#E5E7EB")
        elif r % 2 == 0:
            draw.rectangle([margin, y0, width - margin, y1], fill="#F9FAFB")
        x = margin
        for i, text in enumerate(line):
            f = font_bold if r == 0 else font
            text_w = f.getlength(text)
            tx = x + widths[i] - pad_x - text_w if (COLUMNS[i] in numeric) else x + pad_x
            draw.text((tx, (y0 + y1) / 2), text, font=f, fill=ink, anchor="lm")
            x += widths[i]
    # сетка
    for r in range(len(table) + 1):
        y = top + r * row_h
        draw.line([margin, y, width - margin, y], fill=grid, width=2)
    x = margin
    for w in [0] + widths:
        x += w
        draw.line([x, top, x, top + row_h * len(table)], fill=grid, width=2)
    draw.line([margin, top, margin, top + row_h * len(table)], fill=grid, width=2)

    img.save(path, format="PNG", optimize=True)


# --------------------------------------------------------------------------- запуск
def generate_all(out_dir: Path = DEFAULT_OUT) -> dict[str, Path]:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    rows = build_rows()
    validate_rows(rows)

    paths = {
        "xlsx": out_dir / "transactions_2026.xlsx",
        "pdf": out_dir / "business_metrics.pdf",
        "png": out_dir / "transactions_screenshot.png",
        "expected": out_dir / "expected_metrics.json",
    }
    write_xlsx(rows, paths["xlsx"])
    write_pdf(paths["pdf"])
    write_png(rows, paths["png"])
    with paths["expected"].open("w", encoding="utf-8", newline="\n") as fh:  # LF на любой ОС
        fh.write(json.dumps(compute_expected(rows), ensure_ascii=False, indent=2) + "\n")
    return paths


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT, help="каталог вывода (по умолчанию data/demo)")
    args = parser.parse_args(argv)

    for name, path in generate_all(args.out).items():
        print(f"{name:9} {path}  ({path.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
