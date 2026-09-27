"""Генератор синтетических демо-данных продаж: data/demo/sales_2026.xlsx и документы базы знаний data/demo/sales_docs/*.pdf.

Полностью вымышленная розничная сеть «SalesDemo KZ»; фиксированный seed, файлы воспроизводимы побайтно.

Набор: январь — июнь 2026, 4 региона × 3 канала × 4 категории = 288 строк.
Колонки: Month, Region, Channel, Category, Orders, Cancelled_Orders, Revenue_KZT.
Особенности: Electronics растёт быстрее всех, Sports слегка падает, Mobile отменяет заказы чаще всех, в апреле выручка
Astana проваливается (причина в данных не указана).

    python scripts/generate_sales_demo.py
"""
from __future__ import annotations

import random
import sys
from datetime import datetime
from pathlib import Path

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from scripts.generate_demo_data import FIXED_DATE, _fonts, _normalize_zip, _PdfWriter  # noqa: E402

DEMO_DIR = Path(__file__).resolve().parent.parent / "data" / "demo"
OUT = DEMO_DIR / "sales_2026.xlsx"
DOCS_DIR = DEMO_DIR / "sales_docs"
COLUMNS = ["Month", "Region", "Channel", "Category", "Orders", "Cancelled_Orders", "Revenue_KZT"]
MONTHS = [f"2026-0{m}" for m in range(1, 7)]
REGIONS = {"Almaty": 1.0, "Astana": 0.8, "Shymkent": 0.5, "Karaganda": 0.35}
CHANNELS = {"Web": 1.0, "Mobile": 0.9, "Offline": 0.6}
CATEGORIES = {  # вес, месячный рост заказов, средняя сумма заказа, KZT
    "Electronics": (0.9, 0.06, 45000),
    "Clothing": (1.2, 0.01, 12000),
    "Home": (0.8, 0.03, 22000),
    "Sports": (0.6, -0.01, 9000),
}
CANCEL_RATE = {"Web": 0.04, "Mobile": 0.065, "Offline": 0.02}
BASE_ORDERS = 420
DIP = {("2026-04", "Astana"): 0.62}  # коэффициент выручки в аномальной точке


def build_rows(seed: int = 2026) -> list[dict]:
    rng = random.Random(seed)
    rows = []
    for index, month in enumerate(MONTHS):
        for region, region_weight in REGIONS.items():
            for channel, channel_weight in CHANNELS.items():
                for category, (weight, growth, check) in CATEGORIES.items():
                    orders = round(BASE_ORDERS * region_weight * channel_weight * weight * (1 + growth) ** index * rng.uniform(0.95, 1.05))
                    cancelled = round(orders * CANCEL_RATE[channel] * rng.uniform(0.85, 1.15))
                    revenue = round((orders - cancelled) * check * rng.uniform(0.95, 1.05) * DIP.get((month, region), 1.0))
                    rows.append(
                        {"Month": month, "Region": region, "Channel": channel, "Category": category,
                         "Orders": orders, "Cancelled_Orders": cancelled, "Revenue_KZT": revenue}
                    )
    return rows


def write_xlsx(rows: list[dict], path: Path = OUT) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "sales"
    sheet.append(COLUMNS)
    for row in rows:
        sheet.append([row[c] for c in COLUMNS])
    fill = PatternFill("solid", fgColor="4F46E5")
    for cell in sheet[1]:
        cell.font, cell.fill, cell.alignment = Font(bold=True, color="FFFFFF"), fill, Alignment(horizontal="center")
    for letter, width in zip("ABCDEFG", (10, 12, 10, 14, 10, 18, 14)):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A2"
    workbook.properties.creator = "DataStory AI demo generator"
    workbook.properties.created = workbook.properties.modified = FIXED_DATE
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    _normalize_zip(path)
    return path


# --------------------------------------------------------------------------- документы базы знаний
DOCS: dict[str, tuple[str, str, list[tuple[str, list[str]]]]] = {
    "sales_metrics.pdf": (
        "SalesDemo KZ — определения показателей продаж",
        "Документ создан для демонстрации DataStory AI. Сеть, данные и события в нём полностью вымышлены.",
        [
            ("1. Набор данных sales_2026.xlsx", [
                "Файл sales_2026.xlsx содержит агрегированные продажи за январь — июнь 2026 года: одна строка на сочетание месяца, региона, "
                "канала и категории товаров. Колонки: Month, Region, Channel, Category, Orders, Cancelled_Orders, Revenue_KZT.",
            ]),
            ("2. Выручка (Revenue)", [
                "Выручка (Revenue) — сумма оплаченных заказов в тенге за период, колонка Revenue_KZT. Отменённые заказы в выручку не входят. "
                "Выручку суммируют по любым строкам: по месяцам, регионам, каналам и категориям.",
            ]),
            ("3. Заказы и отмены", [
                "Заказы (Orders) — число оформленных заказов, включая те, что были отменены позже. Отменённые заказы (Cancelled_Orders) — "
                "число заказов, отменённых клиентом или магазином до выдачи.",
            ]),
            ("4. Средний чек (Average Order Value, AOV)", [
                "Средний чек (Average Order Value, AOV) — средняя сумма одного заказа в тенге. Формула: AOV = Revenue_KZT / Orders. "
                "При объединении строк (месяцы, регионы, каналы, категории) сначала суммируют Revenue_KZT и Orders, затем делят одну сумму на другую. "
                "Среднее арифметическое средних чеков по строкам не используется, так как оно не учитывает разный объём строк.",
            ]),
            ("5. Доля отмен (Cancellation Rate)", [
                "Доля отмен (Cancellation Rate) — процент заказов, которые были отменены. Формула: Cancellation Rate = Cancelled_Orders / Orders × 100. "
                "При объединении строк сначала суммируют Cancelled_Orders и Orders, затем делят.",
            ]),
            ("6. Показатели без данных", [
                "Ниже описаны показатели, которые в компании считаются по другим системам. В наборе sales_2026.xlsx нужных колонок нет, поэтому по этим "
                "данным они не рассчитываются. Доля возвратов (Return Rate) — процент проданных товаров, возвращённых покупателем после выдачи. "
                "Конверсия (Conversion Rate) — доля посетителей сайта, оформивших заказ. Пожизненная ценность клиента (Customer Lifetime Value) — "
                "ожидаемая выручка от клиента за всё время сотрудничества.",
            ]),
        ],
    ),
    "sales_channels_regions.pdf": (
        "SalesDemo KZ — каналы, регионы и категории",
        "Документ создан для демонстрации DataStory AI. Сеть, данные и события в нём полностью вымышлены.",
        [
            ("1. Каналы продаж", [
                "Web — интернет-магазин на сайте компании, доставка курьером и в пункты выдачи. Mobile — заказы из мобильного приложения; "
                "в этом канале клиенты отменяют заказы чаще всего. Offline — розничные магазины сети, самовывоз и покупка на месте.",
            ]),
            ("2. Регионы", [
                "Сеть работает в четырёх регионах: Almaty — крупнейший рынок, Astana — второй по объёму, Shymkent и Karaganda — небольшие "
                "региональные рынки с меньшим числом магазинов.",
            ]),
            ("3. Категории товаров", [
                "Electronics — смартфоны, ноутбуки и бытовая электроника с самым высоким средним чеком. Clothing — одежда и обувь с самым большим "
                "числом заказов. Home — товары для дома. Sports — спортивные товары с самым низким средним чеком.",
            ]),
        ],
    ),
    "sales_calendar_2026.pdf": (
        "SalesDemo KZ — календарь событий 2026",
        "Документ создан для демонстрации DataStory AI. Сеть, данные и события в нём полностью вымышлены.",
        [
            ("1. События 2026 года", [
                "В феврале 2026 года запущена программа лояльности для покупателей интернет-магазина. В марте 2026 года в канале Web проходила "
                "промо-неделя со скидками на категорию Home. В апреле 2026 года в Astana проводилась реконструкция торгового центра, где "
                "расположены магазины сети.",
            ]),
            ("2. Как читать события", [
                "Событие совпадает по времени с изменением показателя, но не доказывает причинную связь. Для выводов о причинах нужны "
                "дополнительные данные: журналы работы магазинов, сроки и состав работ.",
            ]),
        ],
    ),
}


def write_pdf(path: Path, title: str, intro: str, sections: list[tuple[str, list[str]]]) -> Path:
    regular, bold, _, _ = _fonts()
    doc = pymupdf.open()
    pdf = _PdfWriter(doc, regular, bold)
    pdf.text(title, font=bold, size=pdf.TITLE, gap_after=6)
    pdf.text(intro, font=regular, size=9.5, gap_after=10, color=(0.42, 0.45, 0.5))
    for heading, paragraphs in sections:
        pdf.heading(heading)
        for paragraph in paragraphs:
            pdf.paragraph(paragraph)
    stamp = "D:20260701000000Z"
    doc.set_metadata({"title": title, "author": "DataStory AI demo generator", "subject": "Синтетические демонстрационные данные",
                      "creator": "scripts/generate_sales_demo.py", "producer": "PyMuPDF", "creationDate": stamp, "modDate": stamp})
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path, garbage=4, deflate=True, no_new_id=True)
    doc.close()
    return path


def generate_all(out_dir: Path = DEMO_DIR) -> dict[str, Path]:
    paths = {"xlsx": write_xlsx(build_rows(), out_dir / "sales_2026.xlsx")}
    for name, (title, intro, sections) in DOCS.items():
        paths[name] = write_pdf(out_dir / "sales_docs" / name, title, intro, sections)
    return paths


if __name__ == "__main__":
    for path in generate_all().values():
        print(path)
